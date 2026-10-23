import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import urlopen

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from fontTools.ttLib import TTFont


ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / ".implementation-work/public-template-assets"
FONT_URL = "https://raw.githubusercontent.com/notofonts/noto-cjk/main/Sans/OTF/SimplifiedChinese/NotoSansCJKsc-Regular.otf"
LICENSE_URL = "https://raw.githubusercontent.com/notofonts/noto-cjk/main/Sans/LICENSE"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def download(url, target, maximum):
    if target.exists():
        return
    with urlopen(url, timeout=120) as response:
        if response.status != 200:
            raise RuntimeError("PUBLIC_ASSET_RESPONSE_INVALID")
        content = response.read(maximum + 1)
    if not content or len(content) > maximum:
        raise RuntimeError("PUBLIC_ASSET_CAPACITY_INVALID")
    target.write_bytes(content)


def main():
    DIRECTORY.mkdir(parents=True, exist_ok=True)
    DIRECTORY.chmod(0o700)
    font_path, license_path = DIRECTORY / "NotoSansCJKsc-Regular.otf", DIRECTORY / "OFL-LICENSE.txt"
    download(FONT_URL, font_path, 32 * 1024 * 1024)
    download(LICENSE_URL, license_path, 32768)
    if "SIL OPEN FONT LICENSE Version 1.1" not in license_path.read_text():
        raise RuntimeError("PUBLIC_FONT_LICENSE_REQUIRED")
    title = "SIMULATION 产品资料整理验收模板"
    with TTFont(font_path, lazy=False) as font:
        family = font["name"].getBestFamilyName()
        if not family or "Noto Sans CJK SC" not in family or not set(map(ord, title)) <= set(font.getBestCmap()):
            raise RuntimeError("PUBLIC_FONT_GLYPHS_REQUIRED")
    layout_path = DIRECTORY / "simulation-product-overview.docx"
    if not layout_path.exists():
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25)
        style = document.styles["Normal"]
        style.font.name, style.font.size = family, Pt(11)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), family)
        document.add_heading(title, 0)
        document.add_paragraph("用途：本地流程验收。版式为自行拟定的 A4 页面与 25mm 页边距，不表示监管要求。")
        document.add_heading("1 资料目录", 1)
        document.add_paragraph("填写资料名称、固定版本、来源链接及核验人员。")
        document.add_heading("2 产品描述", 1)
        document.add_paragraph("保留原始材料中的型号、尺寸、单位和适用条件；未知内容保持待确认。")
        document.add_heading("3 核验记录", 1)
        document.add_paragraph("记录核验范围、决定理由和相关固定版本。")
        document.add_paragraph("字体：Noto Sans CJK SC；许可：SIL Open Font License 1.1；原始字体及许可一并保留。")
        document.core_properties.title = title
        document.core_properties.subject = "SIMULATION 本地模板审查验收"
        document.core_properties.author = "Med-device-Agent Workspace"
        document.save(layout_path)
    parsed = Document(layout_path)
    if parsed.paragraphs[0].text != title or not math.isclose(parsed.sections[0].page_width.mm, 210, abs_tol=0.05):
        raise RuntimeError("PUBLIC_LAYOUT_VALIDATION_FAILED")
    text_path = DIRECTORY / "public-template-notes.json"
    if not text_path.exists():
        value = {"kind": "TEXT", "text": "SIMULATION 本地模板审查依据。使用公开的 Noto Sans CJK SC 原始字体和 SIL Open Font License 1.1。字体来源：" + FONT_URL + "；许可来源：" + LICENSE_URL + "。A4 页面、25mm 页边距和目录编号为本地验收自行拟定的编辑规则。没有使用客户设备事实或专业批准结果。", "marks": []}
        text_path.write_text(json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True), encoding="utf-8")
    license_content_path = DIRECTORY / "ofl-license-text.json"
    if not license_content_path.exists():
        license_content_path.write_text(json.dumps({"kind": "TEXT", "text": license_path.read_text(), "marks": []},
            ensure_ascii=False, separators=(",", ":"), sort_keys=True), encoding="utf-8")
    manifest = {"data_kind": "SIMULATION", "prepared_at": datetime.now(timezone.utc).isoformat(), "font_family": family,
        "font_source": FONT_URL, "license_source": LICENSE_URL, "license": "SIL-OFL-1.1", "layout_origin": "SELF_AUTHORED",
        "assets": [{"kind": kind, "path": str(path.relative_to(ROOT)), "sha256": digest(path), "size_bytes": str(path.stat().st_size), "media_type": media}
            for kind, path, media in (("FONT", font_path, "font/otf"), ("FORMAT", layout_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), ("DEPENDENCY", text_path, "application/json"), ("LICENSE", license_content_path, "application/json"))]}
    (DIRECTORY / "provenance.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    for item in DIRECTORY.iterdir():
        if item.is_file():
            item.chmod(0o600)
    print(json.dumps({"state": "PREPARED", "data_kind": "SIMULATION", "assets": [{"kind": item["kind"], "sha256": item["sha256"]} for item in manifest["assets"]]}))


if __name__ == "__main__":
    main()
