from pathlib import Path
from hashlib import sha256
from importlib.metadata import version
import json
import subprocess
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".lld-data-work"
WORK.mkdir(exist_ok=True)
PUBLIC = ROOT / "ui-delivery/assets/public-evidence"
OFFICIAL = ROOT / "ui-delivery/research/official"
report = {"libraries": {name: version(name) for name in ("python-docx", "pypdf")}, "files": []}
for path in sorted(PUBLIC.glob("*.doc*")) + sorted(OFFICIAL.glob("*.doc")):
    if path.suffix == ".docx":
        document = Document(path)
        blocks = [{"paragraph_index": i, "text": p.text} for i, p in enumerate(document.paragraphs) if p.text]
        tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
        content = {"paragraphs": blocks, "tables": tables}
    else:
        result = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)], check=True, capture_output=True)
        content = {"text": result.stdout.decode("utf-8")}
    report["files"].append({"path": str(path.relative_to(ROOT)), "sha256": sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size, "content": content})
for path in sorted(PUBLIC.glob("*.pdf")):
    reader = PdfReader(path)
    pages = [{"page_index": i, "text": page.extract_text(extraction_mode="plain")} for i, page in enumerate(reader.pages)]
    report["files"].append({"path": str(path.relative_to(ROOT)), "sha256": sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size, "pages": pages})
(WORK / "sample-inspection.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
for item in report["files"]:
    content = item.get("content", {})
    print(json.dumps({"path": item["path"], "sha256": item["sha256"], "bytes": item["bytes"], "paragraph_count": len(content.get("paragraphs", [])), "table_count": len(content.get("tables", [])), "page_count": len(item.get("pages", []))}, ensure_ascii=False))
