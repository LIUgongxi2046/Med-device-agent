from pathlib import Path
from hashlib import sha256
from datetime import datetime, timezone
from importlib.metadata import version
from urllib.parse import unquote, urlparse
import json
import platform
import subprocess
import sys
import graphlib
from docx import Document
from pypdf import PdfReader
from markdown_it import MarkdownIt
from jsonschema import Draft202012Validator, FormatChecker
import sqlglot
from sqlglot import exp

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data-design"
schema = json.loads((OUT / "dataset.schema.json").read_text(encoding="utf-8"))
dataset = json.loads((OUT / "sample_dataset.json").read_text(encoding="utf-8"))
Draft202012Validator.check_schema(schema)
Draft202012Validator(schema, format_checker=FormatChecker()).validate(dataset)
sources = {source["source_id"]: source for source in dataset["sources"]}
assert len(sources) == len(dataset["sources"])
assert len({record["record_id"] for record in dataset["records"]}) == len(dataset["records"])
assert len({record["chunk_id"] for record in dataset["records"]}) == len(dataset["records"])
source_receipts = []
graph = {}
for source_id, source in sources.items():
    path = ROOT / source["source_ref"]
    data = path.read_bytes()
    assert sha256(data).hexdigest() == source["sha256"], source_id
    assert len(data) == source["size_bytes"], source_id
    assert source["production_status"] == "PENDING_APPROVAL", source_id
    assert "APPROVED_PRODUCTION" not in source["allowed_use"], source_id
    parent = source["parent_source_id"]
    graph[source_id] = [parent] if parent else []
    if parent:
        assert parent in sources and parent != source_id
    source_receipts.append({"source_id": source_id, "source_ref": source["source_ref"], "sha256": source["sha256"], "status": "PASS"})
list(graphlib.TopologicalSorter(graph).static_order())

original_receipts = []
provenance = json.loads((ROOT / "ui-delivery/assets/public-evidence/provenance.json").read_text(encoding="utf-8"))
for entry in provenance:
    original = ROOT / "ui-delivery/assets/public-evidence" / entry["original"]
    preview = original.with_suffix(".pdf")
    assert sha256(original.read_bytes()).hexdigest() == entry["originalSHA256"]
    assert sha256(preview.read_bytes()).hexdigest() == entry["previewSHA256"]
    original_receipts.append({"original": str(original.relative_to(ROOT)), "preview": str(preview.relative_to(ROOT)), "status": "PASS"})

text_derivations = []
for source in sources.values():
    if source["format"] != "TXT":
        continue
    parent = sources[source["parent_source_id"]]
    result = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(ROOT / parent["source_ref"])], check=True, capture_output=True)
    assert result.stdout == (ROOT / source["source_ref"]).read_bytes(), source["source_id"]
    text_derivations.append({"source_id": source["source_id"], "original_source_id": parent["source_id"], "status": "PASS"})

record_receipts = []
parser = MarkdownIt()
for record in dataset["records"]:
    source = sources[record["source_id"]]
    assert record["source_version"] == source["sha256"]
    assert not record["formal_evidence_eligible"]
    path = ROOT / source["source_ref"]
    locator = record["locator"]
    kind = locator["kind"]
    if kind == "DOCX":
        document = Document(path)
        address = locator["address"]
        if address["kind"] == "PARAGRAPH":
            actual = document.paragraphs[address["paragraph_index"]].text
        elif address["kind"] == "CELL":
            actual = document.tables[address["table_index"]].cell(address["row_index"], address["column_index"]).text
        else:
            raise ValueError("当前样本尚未包含 DOCX 复合节点定位：" + record["record_id"])
    elif kind == "PDF":
        reader = PdfReader(path)
        assert locator["page_basis"] == "CONVERTED_PREVIEW"
        assert source["nature"] == "DERIVED_PREVIEW"
        actual = reader.pages[locator["page_index"]].extract_text(extraction_mode="plain")
    elif kind == "TXT":
        raw = path.read_bytes().decode("utf-8")
        start, end = locator["char_start"], locator["char_end"]
        assert 0 <= start < end <= len(raw)
        assert raw[:start].count("\n") + 1 == locator["line_start"]
        assert raw[:end].count("\n") + 1 == locator["line_end"]
        actual = raw[start:end]
    elif kind == "MARKDOWN":
        raw = path.read_text(encoding="utf-8")
        token = parser.parse(raw)[locator["node_index"]]
        assert token.map == [locator["line_start"] - 1, locator["line_end"]]
        actual = "".join(raw.splitlines(keepends=True)[locator["line_start"] - 1:locator["line_end"]])
    else:
        raise ValueError("当前样本包含尚未配置核查器的格式：" + kind)
    assert actual == record["text"], record["record_id"]
    assert sha256(actual.encode("utf-8")).hexdigest() == record["quote_hash"]
    record_receipts.append({"record_id": record["record_id"], "format": kind, "quote_hash": record["quote_hash"], "reread": "PASS"})

sql_text = (OUT / "core.sql").read_text(encoding="utf-8")
statements = sqlglot.parse(sql_text, read="mysql")
assert len(statements) == 18
assert all(isinstance(statement, exp.Create) for statement in statements)
table_names = [statement.this.this.name for statement in statements]
assert len(set(table_names)) == len(table_names)
for statement in statements:
    columns = [column.name for column in statement.this.expressions if isinstance(column, exp.ColumnDef)]
    assert "tenant_id" in columns
    for foreign_key in statement.find_all(exp.ForeignKey):
        assert foreign_key.expressions[0].name == "tenant_id"

document_text = (ROOT / "LLD-数据.md").read_text(encoding="utf-8")
tokens = parser.parse(document_text)
links = []
for token in tokens:
    for child in token.children or []:
        if child.type != "link_open":
            continue
        href = child.attrGet("href")
        if urlparse(href).scheme or href.startswith("#"):
            continue
        target = ROOT / unquote(href.split("#")[0])
        if target == OUT / "validation-report.json":
            continue
        assert target.is_file(), href
        links.append(href)
for token in tokens:
    if token.type == "fence" and token.info.strip() == "sql":
        parsed = sqlglot.parse(token.content, read="mysql")
        assert all(isinstance(statement, exp.Select) for statement in parsed)

inputs = [ROOT / "PRD.md", ROOT / "HLD.md", ROOT / "LLD-数据.md", OUT / "core.sql", OUT / "dataset.schema.json", OUT / "sample_dataset.json", OUT / "requirements.txt", Path(__file__), OUT / "build_samples.py", OUT / "inspect_samples.py", ROOT / "ui-delivery/assets/public-evidence/provenance.json"]
report = {
    "checked_at": datetime.now(timezone.utc).isoformat(),
    "capability": "READONLY", "outcome_type": "QUALITY", "status": "PASS",
    "command": ".lld-data-work/venv/bin/python data-design/validate_delivery.py",
    "environment": {"python_executable": sys.executable, "python_version": platform.python_version(), "platform": platform.platform(), "working_directory": str(ROOT), "libraries": {name: version(name) for name in ("python-docx", "pypdf", "markdown-it-py", "jsonschema", "sqlglot")}},
    "summary": {"sources": len(sources), "records": len(record_receipts), "schema": "PASS", "source_reread": "PASS", "locator_reread": "PASS", "original_preview_pairs": len(original_receipts), "text_derivations": len(text_derivations), "mysql_ddl_static_parse": "PASS", "table_count": len(statements), "table_names": table_names, "local_link_count": len(links)},
    "source_receipts": source_receipts, "record_receipts": record_receipts,
    "original_preview_pairs": original_receipts, "text_derivations": text_derivations,
    "inputs": [{"path": str(path.relative_to(ROOT)), "sha256": sha256(path.read_bytes()).hexdigest()} for path in inputs],
    "not_run": ["MYSQL_8_4_EXECUTION", "QDRANT_IMPORT_AND_SEARCH", "MODEL_API", "CUSTOMER_BUSINESS_ACCEPTANCE", "CROSS_TENANT_AUTHORIZATION", "SEVEN_FORMAT_ACCEPTANCE", "OCR", "VISUAL_LAYOUT", "FULL_PRIVACY_SCAN", "PERMISSION_REVOCATION", "UPDATE_APPROVAL", "HISTORY_REPLAY", "DELETION_AND_BACKUP_RESTORE"]
}
(OUT / "validation-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
assert json.loads((OUT / "validation-report.json").read_text(encoding="utf-8"))["summary"] == report["summary"]
print(json.dumps(report["summary"], ensure_ascii=False))
