from pathlib import Path
from hashlib import sha256
from importlib.metadata import version
from datetime import datetime, timezone
import json
from docx import Document
from pypdf import PdfReader
from markdown_it import MarkdownIt
from jsonschema import Draft202012Validator, FormatChecker

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data-design"
PUBLIC = "ui-delivery/assets/public-evidence/"
OFFICIAL = "ui-delivery/research/official/"
BEIJING = "https://www.beijing.gov.cn/zhengce/zhengcefagui/qtwj/"
sources = []
records = []


def register(source_id, path, publisher_version, url, nature, format_name, parent=None):
    data = (ROOT / path).read_bytes()
    sources.append({
        "source_id": source_id, "source_ref": path, "owner": "项目资料所有者" if url is None else "国家药品监督管理局",
        "official_url": url, "publisher_version": publisher_version, "sha256": sha256(data).hexdigest(),
        "size_bytes": len(data), "format": format_name, "nature": nature, "parent_source_id": parent,
        "license": "未取得生产使用和外部分发授权记录；当前按用户要求开展本地资料审查",
        "allowed_use": ["LOCAL_DESIGN_REVIEW", "LOCAL_PARSER_CHECK"],
        "sensitivity": "INTERNAL" if url is None else "PUBLIC", "retention": "待企业批准；禁止自动永久清理",
        "production_status": "PENDING_APPROVAL"
    })
    return sources[-1]


def add(source, text, locator, section, checks, processor):
    quote_hash = sha256(text.encode("utf-8")).hexdigest()
    chunk_hash = sha256((source["sha256"] + processor + json.dumps(locator, sort_keys=True)).encode()).hexdigest()
    records.append({
        "record_id": "sample-" + chunk_hash[:24], "source_id": source["source_id"], "source_version": source["sha256"],
        "parse_version": processor, "chunk_id": chunk_hash, "section_title": section, "text": text,
        "quote_hash": quote_hash, "locator": locator, "checks": checks, "formal_evidence_eligible": False
    })


clinical = register("official-clinical-2021-73-a4", PUBLIC + "clinical-evaluation-2021.docx", "2021年第73号附件4", BEIJING + "202204/t20220408_2669595.html", "OFFICIAL_ORIGINAL", "DOCX")
safety = register("official-safety-2021-121-a9", PUBLIC + "safety-principles-2021.doc", "2021年第121号附件9", BEIJING + "202307/t20230721_3204148.html", "OFFICIAL_ORIGINAL", "DOC_REFERENCE")
for name, source_id, release, url in [
    ("registration-format-2021-121", "official-format-2021-121-a4", "2021年第121号附件4", BEIJING + "202307/t20230721_3204148.html"),
    ("registration-requirements-2021-121", "official-requirements-2021-121-a5", "2021年第121号附件5", BEIJING + "202307/t20230721_3204148.html"),
    ("technical-requirements-2022-8", "official-ptr-2022-8", "2022年第8号附件", "https://mpa.hunan.gov.cn/mpa/yszx/jssp/jsbzx/202202/t20220215_22482327.html")
]:
    original = register(source_id, OFFICIAL + name + ".doc", release, url, "OFFICIAL_ORIGINAL", "DOC_REFERENCE")
    text_source = register(source_id + "-text", OFFICIAL + name + ".txt", release, url, "DERIVED_TEXT", "TXT", source_id)
    raw = (ROOT / text_source["source_ref"]).read_bytes().decode("utf-8")
    snippets = {
        "registration-format-2021-121": ["5.各项申报资料中的申请内容应当具有一致性。"],
        "registration-requirements-2021-121": ["6.1综述", "6.2章节目录"],
        "technical-requirements-2022-8": ["（五）如产品技术要求中的内容引用国家标准、行业标准的，应注明相应标准的编号和年代号。", "（六）产品技术要求编号为相应产品的注册证号（备案号）。拟注册（备案）的产品技术要求编号可留空。"]
    }[name]
    for snippet in snippets:
        start = raw.index(snippet)
        end = start + len(snippet)
        add(text_source, raw[start:end], {"kind": "TXT", "encoding": "UTF-8", "newline_policy": "PRESERVE", "char_start": start, "char_end": end, "line_start": raw[:start].count("\n") + 1, "line_end": raw[:end].count("\n") + 1}, release, ["EXACT_TEXT_RANGE", "DERIVATION_REVIEW_REQUIRED"], "utf8-codepoint-v1")

document = Document(ROOT / clinical["source_ref"])
processor = "python-docx-" + version("python-docx") + ":document-paragraphs-v1"
for index in [10, 159, 160]:
    add(clinical, document.paragraphs[index].text, {"kind": "DOCX", "part": "word/document.xml", "address": {"kind": "PARAGRAPH", "paragraph_index": index}}, "适用范围" if index == 10 else "临床试验概述", ["EXACT_PARAGRAPH", "CONTINUATION_REQUIRED"] if index in [159, 160] else ["EXACT_PARAGRAPH"], processor)
add(clinical, document.tables[1].cell(0, 4).text, {"kind": "DOCX", "part": "word/document.xml", "address": {"kind": "CELL", "table_index": 1, "row_index": 0, "column_index": 4}}, "产品对比表表头", ["EXACT_TABLE_CELL", "MULTILINE_CELL"], processor)
for name, parent in [("clinical-evaluation-2021", clinical), ("safety-principles-2021", safety)]:
    source = register(parent["source_id"] + "-preview", PUBLIC + name + ".pdf", parent["publisher_version"], parent["official_url"], "DERIVED_PREVIEW", "PDF", parent["source_id"])
    page = PdfReader(ROOT / source["source_ref"]).pages[0]
    add(source, page.extract_text(extraction_mode="plain"), {"kind": "PDF", "page_index": 0, "page_basis": "CONVERTED_PREVIEW", "printed_label": None, "region": {"kind": "WHOLE_PAGE"}}, "转换预览第一页", ["EXACT_PAGE_TEXT", "NATIVE_LOCATOR_REQUIRED", "LAYOUT_NOT_VERIFIED"], "pypdf-" + version("pypdf") + ":plain-v1")

spec = register("project-prd-v1.3.0", "PRD.md", "V1.3.0", None, "PROJECT_SPECIFICATION", "MARKDOWN")
raw = (ROOT / "PRD.md").read_text(encoding="utf-8")
tokens = MarkdownIt().parse(raw)
token_index = next(i for i, token in enumerate(tokens) if token.type == "heading_open")
line_start, line_end = tokens[token_index].map
add(spec, "".join(raw.splitlines(keepends=True)[line_start:line_end]), {"kind": "MARKDOWN", "node_index": token_index, "line_start": line_start + 1, "line_end": line_end, "heading_path": [tokens[token_index + 1].content]}, "项目需求文档标题", ["PROJECT_SPECIFICATION_ONLY", "EXACT_SOURCE_LINES"], "markdown-it-py-" + version("markdown-it-py") + ":commonmark-v1")

dataset = {"schema_version": "1.0.0", "dataset_version": "local-evidence-2026-09-28-v1", "created_at": datetime.now(timezone.utc).isoformat(), "sources": sources, "records": records}
schema = json.loads((OUT / "dataset.schema.json").read_text(encoding="utf-8"))
Draft202012Validator.check_schema(schema)
Draft202012Validator(schema, format_checker=FormatChecker()).validate(dataset)
(OUT / "sample_dataset.json").write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"sources": len(sources), "records": len(records), "schema_status": "PASS"}, ensure_ascii=False))
