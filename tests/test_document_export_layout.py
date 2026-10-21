"""Deferred B29 renderer tests. Not run during the requested implementation-only phase."""
from io import BytesIO
from types import SimpleNamespace
import unittest
from docx import Document
from docx.oxml.ns import qn
from medical_document.export import ExportRenderer, ExportFile


class ExportNavigationTest(unittest.TestCase):
    def test_toc_uses_actual_bookmark_reference_fields_and_footer_page(self):
        renderer = ExportRenderer.__new__(ExportRenderer)
        document = Document()
        root = SimpleNamespace(node_id='root')
        section = SimpleNamespace(node_id='12345678-1234-1234-1234-123456789012',
                                  content=SimpleNamespace(kind='SECTION', title='Chapter'))
        plan = SimpleNamespace(layout_profile={'toc_title': 'Contents', 'toc_depth': 2,
                                               'page_number_start': 3})
        renderer._navigation(document, plan, root, {'root': [section], section.node_id: []})
        renderer._anchor(document.add_paragraph('Chapter', style='Heading 1'), section.node_id, 1)
        data = BytesIO(); document.save(data); data.seek(0)
        actual = Document(data)
        instructions = [node.text.strip() for node in actual.element.xpath('.//w:instrText')]
        self.assertEqual(instructions, ['PAGEREF medical_12345678_1234_1234_1234_123456789012 \\h'])
        self.assertEqual(actual.element.xpath('.//w:hyperlink')[0].get(qn('w:anchor')),
                         actual.element.xpath('.//w:bookmarkStart')[0].get(qn('w:name')))
        self.assertEqual([node.text.strip() for node in actual.sections[0].footer._element.xpath('.//w:instrText')], ['PAGE'])
        self.assertEqual(actual.sections[0]._sectPr.find(qn('w:pgNumType')).get(qn('w:start')), '3')
        self.assertEqual(actual.settings.element.find(qn('w:updateFields')).get(qn('w:val')), 'true')

    def test_missing_layout_does_not_invent_channel_rules(self):
        renderer = ExportRenderer.__new__(ExportRenderer)
        document = Document()
        renderer._navigation(document, SimpleNamespace(layout_profile=None), None, {})
        self.assertEqual(document.element.xpath('.//w:instrText'), [])
        self.assertEqual(ExportFile.__dataclass_fields__['visual_check'].default, 'NOT_RUN')
        self.assertEqual(ExportFile.__dataclass_fields__['channel_check'].default, 'UNKNOWN')
