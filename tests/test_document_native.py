import csv
import hashlib
from importlib.metadata import version
import json
from pathlib import Path
import unittest
from uuid import UUID

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
import rfc8785
from docx import Document
from docx.table import Table

from medical_contracts.catalogue import SchemaCatalogue
from medical_contracts.dataset import DatasetBindings
from medical_document import NativeParser, SourceInput


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".implementation-work/document-tests"
VERSION = "31dd0d5d-060a-4ed8-a56a-e6ad9617c60d"
PARSE = "d6f4cc72-8db7-4e80-a13f-119cc7cbe59b"


def source(path, kind, **options):
    parse_version = options.pop("parse_version", PARSE)
    return SourceInput(path, kind, VERSION, parse_version, hashlib.sha256(path.read_bytes()).hexdigest(), **options)


class DocumentNativeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        WORK.mkdir(parents=True, exist_ok=True)
        cls.parser = NativeParser(SchemaCatalogue(ROOT))
        cls.bindings = DatasetBindings(SchemaCatalogue(ROOT))
        cls.dataset = json.loads((ROOT / "data-design/sample_dataset.json").read_bytes())
        cls.csv = WORK / "simulation-tasks.csv"
        with cls.csv.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["data_kind", "task_id", "goal"])
            writer.writerow(["SIMULATION", VERSION, "企业任务\n读取验收"])
            writer.writerow(["SIMULATION", PARSE, "分页与过滤验收"])
        cls.xlsx = WORK / "simulation-tasks.xlsx"
        workbook = Workbook(); sheet = workbook.active; sheet.title = "SIMULATION任务"
        sheet.append(["data_kind", "task_id", "goal"])
        sheet.append(["SIMULATION", VERSION, "任务读取验收"])
        sheet["D2"] = "=LEN(C2)"
        workbook.save(cls.xlsx); workbook.close()
        cls.pptx = WORK / "simulation-tasks.pptx"
        presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
        text.text_frame.text = "SIMULATION 任务读取验收"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(7), Inches(2)).table
        table.cell(0, 0).text = "task_id"; table.cell(0, 1).text = "goal"
        table.cell(1, 0).text = VERSION; table.cell(1, 1).text = "分页与过滤验收"
        presentation.save(cls.pptx)

    def verify_nodes(self, result):
        self.assertGreater(len(result.nodes), 0)
        for node in result.nodes:
            self.assertEqual(node.quote_hash, hashlib.sha256(node.raw_text.encode("utf-8")).hexdigest())
            UUID(node.node_id)
            self.bindings.bind(node.locator.model_dump(mode="json"), "locator")

    def test_official_docx_preserves_published_paragraph_and_cell_locations(self):
        path = ROOT / "ui-delivery/assets/public-evidence/clinical-evaluation-2021.docx"
        result = self.parser.parse(source(path, "DOCX"))
        self.verify_nodes(result)
        records = [item for item in self.dataset["records"] if item["source_id"] == "official-clinical-2021-73-a4"]
        self.assertGreater(len(records), 0)
        for record in records:
            matches = [node for node in result.nodes if node.locator.model_dump(mode="json") == record["locator"]]
            self.assertEqual(len(matches), 1, record["record_id"])
            self.assertEqual(matches[0].raw_text, record["text"])
        self.assertEqual(result.processor_manifest["docling_pipeline"], "SimplePipeline")
        self.assertGreater(result.processor_manifest["docling_items"], 0)

    def test_text_preserves_exact_character_ranges_and_newlines(self):
        path = ROOT / "ui-delivery/research/official/registration-format-2021-121.txt"
        result = self.parser.parse(source(path, "TXT"))
        self.verify_nodes(result)
        text = path.read_bytes().decode("utf-8")
        for node in result.nodes:
            self.assertEqual(text[node.locator.char_start:node.locator.char_end], node.raw_text)
        self.assertEqual("".join(node.raw_text for node in result.nodes), text)

    def test_official_docx_tables_stay_between_their_original_paragraphs(self):
        path = ROOT / "ui-delivery/assets/public-evidence/clinical-evaluation-2021.docx"
        original = Document(path)
        blocks = list(original.iter_inner_content())
        result = self.parser.parse(source(path, "DOCX"))
        paragraph_positions = {
            node.locator.address.paragraph_index: index
            for index, node in enumerate(result.nodes)
            if node.locator.address.kind == "PARAGRAPH"
        }
        paragraph_index = 0
        table_index = 0
        for block_index, block in enumerate(blocks):
            if not isinstance(block, Table):
                paragraph_index += 1
                continue
            cells = [index for index, node in enumerate(result.nodes)
                if node.locator.address.kind == "CELL" and node.locator.address.table_index == table_index]
            self.assertTrue(cells, "登记原件的完整表格需要进入阅读结果")
            previous = [index for index in paragraph_positions if index < paragraph_index]
            following = [index for index in paragraph_positions if index >= paragraph_index]
            if previous:
                self.assertLess(paragraph_positions[max(previous)], min(cells),
                    f"原件主体位置 {block_index} 的表格必须接续原段落")
            if following:
                self.assertLess(max(cells), paragraph_positions[min(following)],
                    f"原件主体位置 {block_index} 的表格必须位于后续段落以前")
            table_index += 1
        self.assertEqual(table_index, 5)
        self.assertEqual(len(original.paragraphs), 175)
        self.verify_nodes(result)

    def test_markdown_keeps_source_lines_and_heading_path(self):
        path = ROOT / "PRD.md"
        result = self.parser.parse(source(path, "MARKDOWN"))
        self.verify_nodes(result)
        lines = path.read_text(encoding="utf-8").splitlines(keepends=True)
        for node in result.nodes:
            self.assertEqual("".join(lines[node.locator.line_start - 1:node.locator.line_end]), node.raw_text)
        self.assertTrue(any(node.locator.heading_path for node in result.nodes))

    def test_csv_quoted_newlines_use_logical_record_indices(self):
        result = self.parser.parse(source(self.csv, "CSV", data_kind="SIMULATION"))
        self.verify_nodes(result)
        goals = [node for node in result.nodes if node.locator.column_indices == [2]]
        self.assertEqual([node.raw_text for node in goals], ["goal", "企业任务\n读取验收", "分页与过滤验收"])
        self.assertEqual([node.locator.record_index for node in goals], [0, 1, 2])
        self.assertEqual(result.processor_manifest["csv_header_records"], 1)

    def test_xlsx_formula_without_cached_result_keeps_formula_and_reports_requirement(self):
        result = self.parser.parse(source(self.xlsx, "XLSX", data_kind="SIMULATION"))
        self.verify_nodes(result)
        formula = [node for node in result.nodes if node.locator.cell_range == "D2"]
        self.assertEqual(len(formula), 1); self.assertEqual(formula[0].raw_text, "=LEN(C2)")
        self.assertEqual(formula[0].locator.value_basis, "FORMULA")
        self.assertEqual(result.status, "PARTIAL_FAILED")
        self.assertTrue(any(issue.code == "FORMULA_RESULT_MISSING" and issue.locator.cell_range == "D2" for issue in result.issues))
        workbook = load_workbook(self.xlsx, data_only=False)
        self.assertEqual(workbook.active["D2"].value, "=LEN(C2)"); workbook.close()

    def test_pptx_keeps_slide_shape_and_table_cell_identity(self):
        result = self.parser.parse(source(self.pptx, "PPTX", data_kind="SIMULATION"))
        self.verify_nodes(result)
        presentation = Presentation(self.pptx)
        table = presentation.slides[0].shapes[1]
        goal = [node for node in result.nodes if node.raw_text == "分页与过滤验收"]
        self.assertEqual(len(goal), 1)
        self.assertEqual(goal[0].locator.slide_id, str(presentation.slides[0].slide_id))
        self.assertEqual(goal[0].locator.shape_id, str(table.shape_id))
        self.assertEqual(goal[0].locator.table_cell.row_index, 1)
        self.assertEqual(goal[0].locator.table_cell.column_index, 1)
        self.assertEqual(result.processor_manifest["docling_pipeline"], "SimplePipeline")

    def test_pdf_conversion_keeps_preview_identity_and_original_page_order(self):
        path = ROOT / "ui-delivery/assets/public-evidence/clinical-evaluation-2021.pdf"
        result = self.parser.parse(source(path, "PDF", page_basis="CONVERTED_PREVIEW"))
        self.verify_nodes(result)
        self.assertEqual(result.processor_manifest["page_count"], 18)
        self.assertEqual([node.locator.page_index for node in result.nodes], list(range(18)))
        self.assertTrue(all(node.locator.page_basis == "CONVERTED_PREVIEW" for node in result.nodes))

    def test_source_digest_changes_fail_before_processing(self):
        path = ROOT / "PRD.md"
        request = SourceInput(path, "MARKDOWN", VERSION, PARSE, "0" * 64)
        with self.assertRaisesRegex(ValueError, "SOURCE_CONTENT_MISMATCH"):
            self.parser.parse(request)

    def test_processor_manifest_binds_actual_versions_and_processing_parameters(self):
        result = self.parser.parse(source(self.csv, "CSV", data_kind="SIMULATION"))
        manifest = result.processor_manifest
        parameters = json.loads(manifest["parameters_json"])
        self.assertEqual(manifest["parameters_sha256"], hashlib.sha256(rfc8785.dumps(parameters)).hexdigest())
        self.assertEqual(parameters["format"], "CSV")
        self.assertEqual(parameters["network_enabled"], False)
        self.assertEqual(parameters["csv_header_record_count"], 1)
        for package in ("docling", "python-docx", "openpyxl", "pandas", "python-pptx", "pypdf", "markdown-it-py"):
            self.assertEqual(manifest["version:" + package], version(package))

    def test_returned_locator_changes_do_not_modify_fixed_result(self):
        result = self.parser.parse(source(self.csv, "CSV", data_kind="SIMULATION"))
        location = result.nodes[0].locator
        location.column_indices.append(99)
        self.assertEqual(result.nodes[0].locator.column_indices, [0])

    def test_reprocessing_preserves_old_result_and_assigns_new_node_identity(self):
        path = ROOT / "PRD.md"
        original = self.parser.parse(source(path, "MARKDOWN"))
        current = self.parser.parse(source(path, "MARKDOWN", parse_version="025a5a77-674d-446d-a33c-b226a124ff35"))
        self.assertEqual([node.raw_text for node in original.nodes], [node.raw_text for node in current.nodes])
        self.assertEqual([node.locator for node in original.nodes], [node.locator for node in current.nodes])
        self.assertTrue(set(node.node_id for node in original.nodes).isdisjoint(node.node_id for node in current.nodes))
        self.assertEqual(original.parse_version, PARSE)


if __name__ == "__main__":
    unittest.main()
