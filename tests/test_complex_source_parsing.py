"""Deferred B26 checks. Unit fixtures are synthetic, not OCR/provider evidence.

Do not invoke the real OCR engine in this suite. Deployment integration checks
require approved installed assets and a separately authorized scanned corpus.
"""
import hashlib
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from uuid import uuid4

from docx import Document
from medical_contracts.catalogue import SchemaCatalogue
from medical_document.native import NativeParser, SourceInput, _ParseContext
from medical_document.complex_parse import ComplexParser
from medical_document.image_locator import bind_image

ROOT = Path(__file__).resolve().parents[1]


class ComplexSourceParsingTest(unittest.TestCase):
    def test_image_locator_is_independent_and_canonical(self):
        locator = {"schema_version": "1.0.0", "kind": "IMAGE", "frame_index": 0,
            "original_sha256": "a" * 64, "width": 100, "height": 100,
            "x0": 0.0, "y0": 0.0, "x1": 1.0, "y1": 1.0,
            "coordinate_system": "ORIGINAL_TOP_LEFT_NORMALIZED"}
        self.assertEqual(bind_image(locator).read().kind, "IMAGE")
        with self.assertRaises(ValueError):
            bind_image(dict(locator, x0=1.0))

    def test_nested_word_cells_are_recovered_with_stable_paths(self):
        parser = NativeParser(SchemaCatalogue(ROOT))
        doc = Document(); outer = doc.add_table(rows=1, cols=1)
        outer.cell(0, 0).text = "SIMULATION outer"
        outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "SIMULATION inner"
        stream = BytesIO(); doc.save(stream)
        request = SourceInput(Path("simulation.docx"), "DOCX", str(uuid4()), str(uuid4()), hashlib.sha256(stream.getvalue()).hexdigest())
        context = _ParseContext(parser.bindings, request)
        parser._docx(stream.getvalue(), context)
        nested = [node for node in context.nodes if node.raw_text == "SIMULATION inner"]
        self.assertEqual(len(nested), 1)
        self.assertEqual(nested[0].locator.address.kind, "NODE")
        self.assertEqual(nested[0].locator.address.structural_path, [0, 0, 0, 0, 0, 0])
        self.assertNotIn("NESTED_TABLE_REQUIRES_VERIFICATION", [issue.code for issue in context.issues])

    def test_release_hash_rejected_before_engine_loading(self):
        with self.assertRaisesRegex(ValueError, "OCR_RELEASE_CHANGED"):
            ComplexParser(b"{}", "0" * 64)

    def test_image_location_uses_original_digest_not_pdf_identity(self):
        box = SimpleNamespace(l=10, t=20, r=50, b=60)
        bbox = SimpleNamespace(to_top_left_origin=lambda **kwargs: box)
        document = SimpleNamespace(pages={1: SimpleNamespace(size=SimpleNamespace(width=100, height=100))})
        context = SimpleNamespace(request=SimpleNamespace(format="IMAGE", sha256="b" * 64))
        locator = ComplexParser.locator(document, 1, bbox, context)
        self.assertEqual(locator["kind"], "IMAGE")
        self.assertEqual(locator["original_sha256"], "b" * 64)
        self.assertEqual(locator["x0"], 0.1)


if __name__ == "__main__":
    unittest.main()
