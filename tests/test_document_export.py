from dataclasses import replace
import hashlib
from io import BytesIO
import json
import os
import shutil
import subprocess
from pathlib import Path
import unittest
from uuid import uuid4

from docx import Document
from PIL import Image
from pypdf import PdfReader
import rfc8785

from medical_contracts.catalogue import SchemaCatalogue
from medical_document.export import ExportInput, ExportRenderer, FixedAsset, FontBinding, FontResource, ConverterBinding


ROOT = Path(__file__).resolve().parents[1]
SOFFICE = Path(os.environ.get('MEDICAL_SOFFICE') or shutil.which('soffice') or shutil.which('libreoffice') or '/usr/bin/soffice')
FONT = Path(os.environ.get('MEDICAL_EXPORT_FONT', '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'))
FONT_FAMILY = os.environ.get('MEDICAL_EXPORT_FONT_FAMILY', 'Noto Sans CJK SC')
SERIF_DIR = Path(os.environ.get('MEDICAL_EXPORT_SERIF_DIR', '/usr/share/fonts/truetype/liberation2'))



def identity():
    return str(uuid4())


class DocumentExportTest(unittest.TestCase):
    def setUp(self):
        self.work = Path(os.environ['MEDICAL_EXPORT_TEST_ROOT']) / identity()
        self.work.mkdir(parents=True)
        self.ids = [identity() for _ in range(20)]
        root, section, paragraph, text, table, table_paragraph, cell_text, formula, image = self.ids[:9]
        first_row, second_row, first_column, second_column = self.ids[9:13]
        self.image_id = self.ids[13]
        self.tenant, self.project = identity(), identity()
        image_bytes = BytesIO()
        Image.new('RGB', (20, 10), color='blue').save(image_bytes, format='PNG')
        self.asset = FixedAsset(self.image_id, hashlib.sha256(image_bytes.getvalue()).hexdigest(), 'image/png', image_bytes.getvalue())
        self.body = {'schema_version': '1.0.0', 'template_version': self.ids[14], 'nodes': [], 'fields': [], 'evidence_links': [], 'fact_usages': []}
        def node(node_id, parent, ordinal, content):
            return {'node_id': node_id, 'parent_id': parent, 'ordinal': ordinal, 'field_ids': [], 'evidence_link_ids': [], 'content': content}
        self.body['nodes'] = [
            node(root, None, 0, {'kind': 'DOCUMENT', 'title': '模拟产品技术要求'}),
            node(section, root, 0, {'kind': 'SECTION', 'title': '型号及参数'}),
            node(paragraph, section, 0, {'kind': 'PARAGRAPH', 'title': ''}),
            node(text, paragraph, 0, {'kind': 'TEXT', 'text': '模拟型号 SIM-01，电压 1.2300 V，温度 37.00 ℃。', 'marks': ['BOLD', 'UNDERLINE']}),
            node(table, section, 1, {'kind': 'TABLE', 'row_ids': [first_row, second_row], 'column_ids': [first_column, second_column], 'caption': '表1 模拟指标', 'cells': [
                {'cell_id': self.ids[15], 'row_id': first_row, 'column_id': first_column, 'row_span': 2, 'column_span': 2, 'node_ids': [table_paragraph]}]}),
            node(table_paragraph, table, 0, {'kind': 'PARAGRAPH', 'title': ''}),
            node(cell_text, table_paragraph, 0, {'kind': 'TEXT', 'text': '模拟合并单元格 0.0500 mm', 'marks': []}),
            node(formula, section, 2, {'kind': 'FORMULA', 'latex': r'E=\frac{a}{b}', 'description': '公式1 模拟计算关系', 'rendered_artifact_id': None}),
            node(image, section, 3, {'kind': 'IMAGE', 'artifact_id': self.image_id, 'sha256': self.asset.sha256, 'alt_text': '模拟结构图', 'caption': '图1 模拟结构'}),
        ]
        self.body['nodes'][3]['field_ids'] = [self.ids[17]]
        self.body['nodes'][3]['evidence_link_ids'] = [self.ids[19]]
        self.body['fields'] = [{'field_id': self.ids[17], 'node_id': text, 'field_definition_id': self.ids[18],
            'original_value': '1.2300', 'decimal_value': '1.2300', 'unit': 'V', 'model': 'SIM-01', 'condition': None,
            'evidence_link_ids': [self.ids[19]]}]
        quote = '1.2300 V'
        self.body['evidence_links'] = [{'evidence_link_id': self.ids[19], 'node_ids': [text], 'relation': 'SUPPORT', 'evidence': {
            'schema_version': '1.0.0', 'tenant_id': self.tenant, 'project_id': self.project, 'source_id': identity(),
            'source_version': identity(), 'parse_version': identity(), 'chunk_id': identity(), 'locator_kind': 'TXT',
            'locator': {'kind': 'TXT', 'encoding': 'UTF-8', 'newline_policy': 'PRESERVE', 'char_start': 0,
                'char_end': len(quote), 'line_start': 1, 'line_end': 1},
            'quote_hash': hashlib.sha256(quote.encode('utf-8')).hexdigest(), 'approval_revision': identity()}}]
        self.body['fact_usages'] = [{'fact_version_id': identity(), 'node_ids': [text], 'field_ids': [self.ids[17]]}]
        (self.work / 'simulation-source.txt').write_text(quote, encoding='utf-8')
        self.bytes = rfc8785.dumps(self.body)
        template = BytesIO()
        Document().save(template)
        self.template = template.getvalue()
        font_directory = SERIF_DIR
        additional = tuple(FontResource(path, hashlib.sha256(path.read_bytes()).hexdigest(), 'Liberation Serif')
            for path in (font_directory / 'LiberationSerif-Regular.ttf', font_directory / 'LiberationSerif-Italic.ttf'))
        self.font = FontBinding(identity(), FONT, hashlib.sha256(FONT.read_bytes()).hexdigest(), FONT_FAMILY, additional)
        version = subprocess.run([str(SOFFICE), '--version'], check=True, capture_output=True, text=True, timeout=30).stdout.strip()
        if not version.startswith('LibreOffice'):
            raise ValueError('REAL_LIBREOFFICE_VERSION_REQUIRED')
        self.converter = ConverterBinding(identity(), SOFFICE, version, 60,
            hashlib.sha256(SOFFICE.read_bytes()).hexdigest())
        self.plan = ExportInput(self.tenant, self.project, identity(), identity(), self.body['template_version'],
            hashlib.sha256(self.bytes).hexdigest(), hashlib.sha256(self.template).hexdigest(), 'FORMAL_BODY', 'SIMULATION', self.font, self.converter)
        self.renderer = ExportRenderer(SchemaCatalogue(ROOT))
        (self.work / 'fixed-body.json').write_bytes(self.bytes)
        (self.work / 'fixed-template.docx').write_bytes(self.template)

    def test_docx_preserves_complete_fixed_body_and_rejects_overwrite(self):
        result = self.renderer.docx(self.plan, self.bytes, self.template, (self.asset,), self.work / 'docx')
        actual = Document(result.path)
        text = '\n'.join(paragraph.text for paragraph in actual.paragraphs)
        self.assertIn('SIMULATION', text)
        self.assertIn('模拟型号 SIM-01，电压 1.2300 V，温度 37.00 ℃。', text)
        self.assertEqual('模拟合并单元格 0.0500 mm', actual.tables[0].cell(0, 0).text.strip())
        self.assertEqual(1, len(actual.inline_shapes))
        self.assertEqual(1, len(actual.element.xpath('.//m:oMath')))
        self.assertEqual(9, len(result.verified_node_ids))
        self.assertEqual(hashlib.sha256(result.path.read_bytes()).hexdigest(), result.sha256)
        self.assertTrue(any(run.bold and run.underline for paragraph in actual.paragraphs for run in paragraph.runs))
        original = result.path.read_bytes()
        with self.assertRaises(FileExistsError):
            self.renderer.docx(self.plan, self.bytes, self.template, (self.asset,), self.work / 'docx')
        self.assertEqual(original, result.path.read_bytes())

    def test_pdf_uses_actual_converter_and_preserves_text_and_original_docx(self):
        document = self.renderer.docx(self.plan, self.bytes, self.template, (self.asset,), self.work / 'docx')
        original = document.path.read_bytes()
        pdf = self.renderer.pdf(self.plan, document, self.work / 'pdf')
        actual = PdfReader(pdf.path)
        text = '\n'.join(page.extract_text() for page in actual.pages)
        self.assertIn('SIMULATION', text)
        self.assertIn('1.2300', text)
        self.assertIn('0.0500', text)
        self.assertIn('37.00', text)
        self.assertIn('型号及参数', text)
        compact = ''.join(text.split())
        self.assertIn('图1模拟结构', compact)
        self.assertIn('公式1模拟计算关系', compact)
        self.assertEqual(original, document.path.read_bytes())
        self.assertEqual(document.input_sha256, pdf.input_sha256)
        self.assertEqual(document.verified_node_ids, pdf.verified_node_ids)
        with self.assertRaises(ValueError):
            self.renderer.pdf(replace(self.plan, document_revision=identity()), document, self.work / 'other-document-pdf')
        self.assertFalse((self.work / 'other-document-pdf').exists())

    def test_internal_markdown_retains_structure_formula_and_separate_projections(self):
        internal = replace(self.plan, purpose='INTERNAL_EXCHANGE')
        result = self.renderer.markdown(internal, self.bytes, (self.asset,), self.work / 'markdown')
        text = result.path.read_text(encoding='utf-8')
        self.assertIn('SIMULATION', text)
        self.assertIn('1.2300 V', text)
        self.assertIn(r'E=\frac{a}{b}', text)
        self.assertIn('rowspan="2"', text)
        self.assertIn('colspan="2"', text)
        self.assertEqual(self.asset.content, (result.path.parent / (self.asset.artifact_id + '.png')).read_bytes())
        self.assertTrue((result.path.parent / 'internal-projections.json').is_file())
        projections = json.loads((result.path.parent / 'internal-projections.json').read_bytes())
        self.assertEqual(self.body['fields'], projections['fields'])
        self.assertEqual(self.body['evidence_links'], projections['evidence_links'])
        self.assertEqual(self.body['fact_usages'], projections['fact_usages'])
        self.assertEqual(9, len(result.verified_node_ids))
        with self.assertRaises(ValueError):
            self.renderer.markdown(self.plan, self.bytes, (self.asset,), self.work / 'formal-markdown')

    def test_input_asset_font_and_structure_changes_fail_before_output(self):
        for changed in (
            replace(self.plan, body_sha256='0' * 64),
            replace(self.plan, template_sha256='0' * 64),
            replace(self.plan, font=replace(self.font, sha256='0' * 64)),
        ):
            target = self.work / identity()
            with self.assertRaises(ValueError):
                self.renderer.docx(changed, self.bytes, self.template, (self.asset,), target)
            self.assertFalse(target.exists())
        target = self.work / identity()
        with self.assertRaises(ValueError):
            self.renderer.docx(self.plan, self.bytes, self.template, (replace(self.asset, content=b'changed'),), target)
        self.assertFalse(target.exists())
        for changed in (replace(self.plan, tenant_id=identity()), replace(self.plan, project_id=identity())):
            target = self.work / identity()
            with self.assertRaises(ValueError):
                self.renderer.docx(changed, self.bytes, self.template, (self.asset,), target)
            self.assertFalse(target.exists())
        incomplete = dict(self.body)
        incomplete['nodes'] = self.body['nodes'][:-1] + [dict(self.body['nodes'][-1], parent_id=identity())]
        content = rfc8785.dumps(incomplete)
        with self.assertRaises(ValueError):
            self.renderer.docx(replace(self.plan, body_sha256=hashlib.sha256(content).hexdigest()), content, self.template, (self.asset,), target)
        self.assertFalse(target.exists())


if __name__ == '__main__':
    unittest.main()
