from dataclasses import dataclass
from datetime import date, datetime, time
import hashlib
from importlib.metadata import version
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from types import MappingProxyType
from typing import Mapping
from uuid import UUID, uuid5
from zipfile import ZipFile
import rfc8785

from defusedxml import ElementTree
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from openpyxl import load_workbook
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from medical_contracts.catalogue import SchemaCatalogue
from medical_contracts.dataset import DatasetBindings, Locator
from medical_contracts.typed import TypedPayload
from .scan import ClamAVScanner, MAX_FILE_BYTES, NativeScanResult


MAX_EXPANDED_BYTES = 512 * 1024 * 1024
MAX_NODES = 100000
MAX_TEXT_CHARACTERS = 16000000
EXTENSIONS = {"DOCX": {".docx"}, "PDF": {".pdf"}, "XLSX": {".xlsx"}, "CSV": {".csv"},
    "TXT": {".txt"}, "MARKDOWN": {".md", ".markdown"}, "PPTX": {".pptx"}, "HTML": {".html", ".htm"}, "IMAGE": {".image", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}}


@dataclass(frozen=True)
class SourceInput:
    path: Path
    format: str
    source_version: str
    parse_version: str
    sha256: str
    data_kind: str = "REGISTERED"
    page_basis: str = "ORIGINAL"


@dataclass(frozen=True)
class ParsedNode:
    node_id: str
    raw_text: str
    quote_hash: str
    heading_path: tuple[str, ...]
    _locator: TypedPayload[Locator]

    @property
    def locator(self) -> Locator:
        return self._locator.read()


@dataclass(frozen=True)
class ParseIssue:
    code: str
    _locator: TypedPayload[Locator] | None

    @property
    def locator(self) -> Locator | None:
        return None if self._locator is None else self._locator.read()


@dataclass(frozen=True)
class ParsedDocument:
    source_version: str
    parse_version: str
    original_sha256: str
    status: str
    nodes: tuple[ParsedNode, ...]
    issues: tuple[ParseIssue, ...]
    processor_manifest: Mapping[str, object]


@dataclass(frozen=True)
class ScannedDocument:
    document: ParsedDocument
    scan: NativeScanResult


class NativeParser:
    def __init__(self, catalogue: SchemaCatalogue, complex_parser=None):
        self.complex_parser = complex_parser
        self.bindings = DatasetBindings(catalogue)

    def process_command(self, command_path: Path, result_path: Path) -> None:
        from .processing_command import execute
        execute(command_path, result_path, self)

    def scan_and_parse(self, request: SourceInput, release: Mapping[str, object], work_directory: Path) -> ScannedDocument:
        content = self._source_content(request)
        scanner = ClamAVScanner(release)
        result = scanner.scan(content, request.sha256, request.path.suffix.lower(), work_directory)
        return ScannedDocument(self.parse(request), result)

    @staticmethod
    def _source_content(request: SourceInput) -> bytes:
        if request.format not in EXTENSIONS or request.path.suffix.lower() not in EXTENSIONS[request.format]:
            raise ValueError("SOURCE_FORMAT_UNSUPPORTED")
        if request.data_kind not in ("REGISTERED", "SIMULATION") or request.page_basis not in ("ORIGINAL", "CONVERTED_PREVIEW"):
            raise ValueError("SOURCE_PROCESSING_PARAMETERS_INVALID")
        for identifier in (request.source_version, request.parse_version):
            if str(UUID(identifier)) != identifier:
                raise ValueError("SOURCE_VERSION_INVALID")
        size = request.path.stat().st_size
        if size <= 0 or size > MAX_FILE_BYTES or request.path.is_symlink():
            raise ValueError("SOURCE_FILE_LIMIT")
        content = request.path.read_bytes()
        if len(content) != size or hashlib.sha256(content).hexdigest() != request.sha256:
            raise ValueError("SOURCE_CONTENT_MISMATCH")
        return content

    def parse(self, request: SourceInput) -> ParsedDocument:
        content = self._source_content(request)
        context = _ParseContext(self.bindings, request)
        if request.format in ("DOCX", "XLSX", "PPTX"):
            self._package(content)
        if request.format in ("DOCX", "PPTX"):
            converted = self._office(content, request.path.name)
            context.manifest["docling_pipeline"] = "SimplePipeline"
            context.manifest["docling_version"] = version("docling")
            context.manifest["docling_items"] = sum(1 for _ in converted.document.iterate_items())
        from .html_native import parse_html
        handlers = {"HTML": parse_html,"DOCX": self._docx, "PDF": self._pdf, "XLSX": self._xlsx, "CSV": self._csv,
            "TXT": self._text, "MARKDOWN": self._markdown, "PPTX": self._pptx}
        if self.complex_parser is not None:
            context.manifest["complex_engine_release_sha256"] = self.complex_parser.release_sha256
        if request.format == "IMAGE":
            context.manifest["image_locator_schema_sha256"] = "b93cead506584a3899be6dd5a3f7a66420291f67e00d796112620e7e24f77d5c"
            if self.complex_parser is None:
                context.issue("OCR_ENGINE_NOT_CONFIGURED", None)
            else:
                self.complex_parser.parse_pdf(content, context)
        elif request.format == "PDF" and self.complex_parser is not None:
            self.complex_parser.parse_pdf(content, context)
        else:
            handlers[request.format](content, context)
        if not context.nodes:
            context.issue("NATIVE_CONTENT_MISSING", None)
        status = "COMPLETE" if not context.issues else "PARTIAL_FAILED" if context.nodes else "FAILED"
        return ParsedDocument(request.source_version, request.parse_version, request.sha256, status,
            tuple(context.nodes), tuple(context.issues), MappingProxyType(dict(context.manifest)))

    @staticmethod
    def _office(content: bytes, name: str):
        from docling.datamodel.base_models import ConversionStatus, InputFormat
        from docling.datamodel.pipeline_options import ConvertPipelineOptions
        from docling.document_converter import DocumentConverter, PowerpointFormatOption, WordFormatOption
        from docling.pipeline.simple_pipeline import SimplePipeline
        from docling_core.types.io import DocumentStream

        # 完成原件扫描后加载Office转换组件，转换仍关闭外部服务与插件。
        options = ConvertPipelineOptions(enable_remote_services=False, allow_external_plugins=False,
            document_timeout=120, do_picture_classification=False, do_picture_description=False, do_chart_extraction=False)
        converter = DocumentConverter(allowed_formats=[InputFormat.DOCX, InputFormat.PPTX], format_options={
            InputFormat.DOCX: WordFormatOption(pipeline_cls=SimplePipeline, pipeline_options=options),
            InputFormat.PPTX: PowerpointFormatOption(pipeline_cls=SimplePipeline, pipeline_options=options)})
        converted = converter.convert(DocumentStream(name=name, stream=BytesIO(content)),
            raises_on_error=True, max_num_pages=1000, max_file_size=MAX_FILE_BYTES)
        if converted.status != ConversionStatus.SUCCESS:
            raise ValueError("DOCUMENT_CONVERSION_INCOMPLETE")
        return converted

    @staticmethod
    def _package(content: bytes):
        with ZipFile(BytesIO(content)) as package:
            entries = package.infolist()
            if len(entries) > 10000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("OFFICE_PACKAGE_LIMIT")
            for item in entries:
                address = PurePosixPath(item.filename)
                if address.is_absolute() or ".." in address.parts or item.flag_bits & 1:
                    raise ValueError("OFFICE_PACKAGE_INVALID")
                if "vba" in item.filename.lower() or "activex" in item.filename.lower():
                    raise ValueError("OFFICE_ACTIVE_CONTENT_DENIED")
                if item.filename.endswith(".rels"):
                    relationships = ElementTree.fromstring(package.read(item))
                    for relation in relationships:
                        if relation.attrib.get("TargetMode") == "External" and not relation.attrib.get("Type", "").endswith("/hyperlink"):
                            raise ValueError("OFFICE_EXTERNAL_CONTENT_DENIED")

    @classmethod
    def verify_format(cls, content: bytes, suffix: str) -> str:
        formats = [name for name, extensions in EXTENSIONS.items() if suffix.lower() in extensions]
        if len(formats) != 1 or not 0 < len(content) <= MAX_FILE_BYTES:
            raise ValueError("SOURCE_FORMAT_UNSUPPORTED")
        format_name = formats[0]
        if format_name in ("DOCX", "XLSX", "PPTX"):
            cls._package(content)
        if format_name == "DOCX":
            document = Document(BytesIO(content))
            if len(document.paragraphs) + len(document.tables) > MAX_NODES:
                raise ValueError("DOCUMENT_NODE_LIMIT")
        elif format_name == "PDF":
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ValueError("PDF_ENCRYPTED")
            if not 0 < len(reader.pages) <= 1000:
                raise ValueError("PDF_PAGE_LIMIT")
            for page in reader.pages:
                page.get_contents()
        elif format_name == "XLSX":
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                count = 0
                for worksheet in workbook:
                    for row in worksheet.iter_rows():
                        count += len(row)
                        if count > MAX_NODES:
                            raise ValueError("XLSX_CELL_LIMIT")
            finally:
                workbook.close()
        elif format_name == "HTML":
            from .html_native import tree
            tree(content)
        elif format_name == "IMAGE":
            from PIL import Image
            with Image.open(BytesIO(content)) as image:
                expected = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".tif": "TIFF", ".tiff": "TIFF"}
                if suffix.lower() in expected and image.format != expected[suffix.lower()]:
                    raise ValueError("IMAGE_EXTENSION_MISMATCH")
                if image.format not in ("PNG", "JPEG", "TIFF") or getattr(image, "n_frames", 1) != 1:
                    raise ValueError("IMAGE_FORMAT_INVALID")
                if image.width * image.height > 100000000:
                    raise ValueError("IMAGE_PIXEL_LIMIT")
                image.verify()
        elif format_name == "PPTX":
            presentation = Presentation(BytesIO(content))
            if len(presentation.slides) > MAX_NODES:
                raise ValueError("PPTX_SLIDE_LIMIT")
            for slide in presentation.slides:
                if len(slide.shapes) > MAX_NODES:
                    raise ValueError("PPTX_SHAPE_LIMIT")
        else:
            text = content.decode("utf-8", errors="strict")
            if len(text) > MAX_TEXT_CHARACTERS:
                raise ValueError("SOURCE_TEXT_LIMIT")
            if format_name == "CSV":
                table = pd.read_csv(StringIO(text), header=None, dtype=str, keep_default_na=False, na_filter=False,
                    sep=",", quotechar='"', doublequote=True, skip_blank_lines=False, on_bad_lines="error", engine="python")
                if table.size > MAX_NODES:
                    raise ValueError("CSV_CELL_LIMIT")
            elif format_name == "MARKDOWN":
                MarkdownIt("commonmark", {"html": False}).parse(text)
        return format_name

    def _docx(self, content: bytes, context):
        document = Document(BytesIO(content))
        part = str(document.part.partname).lstrip("/")
        headings = []
        paragraph_index = 0
        table_index = 0
        # 阅读顺序采用原件主体顺序，定位继续保留独立的段落与表格编号。
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                style = block.style.name
                if style.startswith("Heading ") and style[8:].isdigit():
                    level = int(style[8:]); headings = headings[:level - 1]; headings.append(block.text)
                context.node(block.text, {"kind": "DOCX", "part": part,
                    "address": {"kind": "PARAGRAPH", "paragraph_index": paragraph_index}}, tuple(headings))
                paragraph_index += 1
            elif isinstance(block, Table):
                self._word_table(block, part, table_index, context, tuple(headings))
                table_index += 1
            else:
                raise TypeError("DOCX_BLOCK_TYPE_INVALID")
        for index, picture in enumerate(document.inline_shapes):
            locator = {"kind": "DOCX", "part": part,
                "address": {"kind": "NODE", "node_id": "inline-picture-" + str(index),
                    "node_hash": hashlib.sha256(picture._inline.xml.encode("utf-8")).hexdigest(), "structural_path": [2, index]}}
            if self.complex_parser is None:
                context.issue("IMAGE_REQUIRES_VERIFICATION", locator)
            else:
                blips = picture._inline.xpath(".//a:blip")
                if len(blips) != 1 or blips[0].embed is None:
                    context.issue("IMAGE_RELATIONSHIP_UNKNOWN", locator)
                else:
                    image = document.part.related_parts[blips[0].embed]
                    self._image(image.blob, Path(str(image.partname)).name, locator, context)

    def _word_table(self, table, part, table_index, context, headings, path=()):
        seen = set()
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells, start=row.grid_cols_before):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                location = {"kind": "DOCX", "part": part, "address": {"kind": "CELL", "table_index": table_index,
                    "row_index": row_index, "column_index": column_index}}
                cell_path = (*path, table_index, row_index, column_index)
                if path:
                    location = {"kind": "DOCX", "part": part, "address": {"kind": "NODE",
                        "node_id": "nested-cell-" + "-".join(map(str, cell_path)),
                        "node_hash": hashlib.sha256(cell._tc.xml.encode("utf-8")).hexdigest(),
                        "structural_path": list(cell_path)}}
                context.node(cell.text, location, headings)
                context.manifest.setdefault("table_structure", []).append({"locator": location,
                    "grid_span": cell._tc.grid_span, "structural_path": list(cell_path)})
                for nested_index, nested in enumerate(cell.tables):
                    if len(cell_path) > 48:
                        raise ValueError("NESTED_TABLE_DEPTH_LIMIT")
                    self._word_table(nested, part, nested_index, context, headings, cell_path)

    @staticmethod
    def _pdf(content: bytes, context):
        reader = PdfReader(BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise ValueError("PDF_ENCRYPTED")
        if len(reader.pages) > 1000:
            raise ValueError("PDF_PAGE_LIMIT")
        labels = reader.page_labels
        context.manifest["page_count"] = len(reader.pages)
        context.manifest["pdf_version"] = version("pypdf")
        for index, page in enumerate(reader.pages):
            locator = {"kind": "PDF", "page_index": index, "page_basis": context.request.page_basis,
                "printed_label": labels[index], "region": {"kind": "WHOLE_PAGE"}}
            text = page.extract_text(extraction_mode="plain")
            if text:
                context.node(text, locator)
            else:
                context.issue("OCR_REQUIRED", locator)
            if len(page.images) > 0:
                context.issue("IMAGE_REQUIRES_VERIFICATION", locator)

    @staticmethod
    def _text(content: bytes, context):
        text = content.decode("utf-8", errors="strict")
        offset = 0
        for line, value in enumerate(text.splitlines(keepends=True), start=1):
            context.node(value, {"kind": "TXT", "encoding": "UTF-8", "newline_policy": "PRESERVE",
                "char_start": offset, "char_end": offset + len(value), "line_start": line, "line_end": line})
            offset += len(value)
        context.manifest["encoding"] = "UTF-8"
        context.manifest["newline_policy"] = "PRESERVE"

    @staticmethod
    def _markdown(content: bytes, context):
        text = content.decode("utf-8", errors="strict")
        lines = text.splitlines(keepends=True)
        tokens = MarkdownIt("commonmark", {"html": False}).parse(text)
        headings = []
        for index, token in enumerate(tokens):
            if token.type == "heading_open":
                level = int(token.tag[1:]); headings = headings[:level - 1]; headings.append(tokens[index + 1].content)
            if token.map is None or token.nesting < 0 or token.level != 0:
                continue
            first, end = token.map
            context.node("".join(lines[first:end]), {"kind": "MARKDOWN", "node_index": index,
                "line_start": first + 1, "line_end": end, "heading_path": list(headings)}, tuple(headings))
        context.manifest["markdown_it_version"] = version("markdown-it-py")
        context.manifest["html_execution"] = "DISABLED"

    @staticmethod
    def _csv(content: bytes, context):
        text = content.decode("utf-8", errors="strict")
        table = pd.read_csv(StringIO(text), header=None, dtype=str, keep_default_na=False, na_filter=False,
            sep=",", quotechar='"', doublequote=True, skip_blank_lines=False, on_bad_lines="error", engine="python")
        if table.size > MAX_NODES:
            raise ValueError("CSV_CELL_LIMIT")
        context.manifest["csv_header_records"] = 1
        context.manifest["csv_dialect"] = "UTF8-COMMA-DOUBLEQUOTE-1"
        for row_index, values in enumerate(table.itertuples(index=False, name=None)):
            for column_index, value in enumerate(values):
                if not isinstance(value, str):
                    raise ValueError("CSV_RECORD_INCOMPLETE")
                context.node(value, {"kind": "CSV", "table_id": context.request.source_version,
                    "dialect_version": "UTF8-COMMA-DOUBLEQUOTE-1", "record_index": row_index,
                    "column_indices": [column_index], "header_record_count": 1})

    @staticmethod
    def _xlsx(content: bytes, context):
        with ZipFile(BytesIO(content)) as package:
            definition = ElementTree.fromstring(package.read("xl/workbook.xml"))
            sheet_ids = {item.attrib["name"]: item.attrib["sheetId"] for item in
                definition.findall("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets/{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet")}
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in formulas:
                if sheet.max_row * sheet.max_column > MAX_NODES:
                    raise ValueError("XLSX_CELL_LIMIT")
                for row in sheet:
                    for cell in row:
                        if cell.value is None:
                            continue
                        locator = {"kind": "XLSX", "sheet_id": sheet_ids[sheet.title], "sheet_name": sheet.title,
                            "cell_range": cell.coordinate, "value_basis": "FORMULA" if cell.data_type == "f" else "LITERAL",
                            "calculation_artifact_id": None}
                        value = cell.value.isoformat() if isinstance(cell.value, (datetime, date, time)) else str(cell.value)
                        context.node(value, locator)
                        if cell.data_type == "f":
                            result = cached[sheet.title][cell.coordinate].value
                            if result is None:
                                context.issue("FORMULA_RESULT_MISSING", locator)
                            else:
                                context.node(str(result), dict(locator, value_basis="CACHED_RESULT"))
        finally:
            formulas.close(); cached.close()
        context.manifest["formula_execution"] = "DISABLED"

    def _pptx(self, content: bytes, context):
        presentation = Presentation(BytesIO(content))
        context.manifest["slide_count"] = len(presentation.slides)
        for slide_index, slide in enumerate(presentation.slides):
            self._slide_shapes(slide.shapes, str(slide.slide_id), slide_index, context)

    def _slide_shapes(self, shapes, slide_id, slide_index, context):
        for shape in shapes:
            locator = {"kind": "PPTX", "slide_id": slide_id, "slide_index": slide_index,
                "shape_id": str(shape.shape_id), "paragraph_index": None, "table_cell": None}
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                self._slide_shapes(shape.shapes, slide_id, slide_index, context)
            if shape.has_text_frame:
                for index, paragraph in enumerate(shape.text_frame.paragraphs):
                    context.node(paragraph.text, dict(locator, paragraph_index=index))
            if shape.has_table:
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        context.node(cell.text, dict(locator, table_cell={"row_index": row_index, "column_index": column_index}))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if self.complex_parser is None:
                    context.issue("IMAGE_REQUIRES_VERIFICATION", locator)
                else:
                    self._image(shape.image.blob, "picture." + shape.image.ext, locator, context)
            elif shape.has_chart:
                context.issue("CHART_REQUIRES_VERIFICATION", locator)

    def _image(self, content, name, locator, context):
        document = self.complex_parser.convert(content, name)
        parts = [item.text for item, _ in document.iterate_items() if getattr(item, "text", None)]
        for table in document.tables:
            parts.extend(cell.text for cell in table.data.table_cells if cell.text)
        if not parts:
            context.issue("IMAGE_TEXT_UNKNOWN", locator)
            return
        context.node("\n".join(parts), locator)
        context.manifest.setdefault("embedded_image_provenance", []).append({
            "node_id": context.nodes[-1].node_id, "image_sha256": hashlib.sha256(content).hexdigest(),
            "engine_release_sha256": self.complex_parser.release_sha256})


class _ParseContext:
    def __init__(self, bindings, request):
        self.bindings = bindings; self.request = request
        self.nodes = []; self.issues = []; self.identities = set(); self.characters = 0
        parameters = {"format": request.format, "page_basis": request.page_basis, "network_enabled": False,
            "csv_header_record_count": 1, "csv_dialect": "UTF8-COMMA-DOUBLEQUOTE-1",
            "encoding": "UTF-8", "newline_policy": "PRESERVE", "formula_execution": "DISABLED",
            "html_execution": "DISABLED", "docx_reading_order": "BODY_CONTENT",
            "document_timeout_seconds": 120, "max_pages": 1000,
            "max_file_bytes": MAX_FILE_BYTES, "max_expanded_bytes": MAX_EXPANDED_BYTES,
            "max_nodes": MAX_NODES, "max_text_characters": MAX_TEXT_CHARACTERS}
        canonical_parameters = rfc8785.dumps(parameters)
        self.manifest = {"format": request.format, "data_kind": request.data_kind,
            "processor_revision": "DOCUMENT-NATIVE-PARSE-2", "schema_version": "1.0.0",
            "locator_schema_sha256": bindings.catalogue.digests["dataset"], "network_enabled": False,
            "parameters_json": canonical_parameters.decode("utf-8"),
            "parameters_sha256": hashlib.sha256(canonical_parameters).hexdigest(), "language": "und"}
        for package in ("docling", "docling-parse", "python-docx", "openpyxl", "pandas", "python-pptx",
                "pypdf", "markdown-it-py", "defusedxml", "medical-contracts", "numpy", "opencv-python"):
            self.manifest["version:" + package] = version(package)

    def node(self, text, locator, headings=()):
        if text == "":
            return
        if not isinstance(text, str):
            raise TypeError("PARSED_TEXT_TYPE_INVALID")
        self.characters += len(text)
        if len(self.nodes) >= MAX_NODES or self.characters > MAX_TEXT_CHARACTERS:
            raise ValueError("PARSE_CONTENT_LIMIT")
        if locator.get("kind") in {"IMAGE", "HTML"}:
            from medical_contracts.source_evidence import SourceEvidenceBindings
            bound = SourceEvidenceBindings(self.bindings.catalogue).bind(locator, "locator")
        else:
            bound = self.bindings.bind(locator, "locator")
        # 新解析版本形成独立节点身份，历史原文和位置保持固定。
        identifier = str(uuid5(UUID(self.request.parse_version), self.request.source_version + ":" + bound.canonical_bytes.decode("utf-8")))
        if identifier in self.identities:
            raise ValueError("PARSE_NODE_DUPLICATE")
        self.identities.add(identifier)
        self.nodes.append(ParsedNode(identifier, text, hashlib.sha256(text.encode("utf-8")).hexdigest(), tuple(headings), bound))

    def issue(self, code, locator):
        if locator is None:
            bound = None
        elif locator.get("kind") in {"IMAGE", "HTML"}:
            from medical_contracts.source_evidence import SourceEvidenceBindings
            bound = SourceEvidenceBindings(self.bindings.catalogue).bind(locator, "locator")
        else:
            bound = self.bindings.bind(locator, "locator")
        self.issues.append(ParseIssue(code, bound))
