from collections import defaultdict
from dataclasses import dataclass
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import signal
import subprocess
from typing import Literal

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from fontTools.ttLib import TTFont
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree
from mathml2omml import convert as mathml_to_omml
from PIL import Image
from pypdf import PdfReader
import rfc8785

from medical_contracts.catalogue import SchemaCatalogue
from medical_contracts.interfaces import DocumentBody, InterfaceBindings
from medical_contracts.wire import uuid_value


SIMULATION_LABEL = 'SIMULATION · 仅用于开发检查'
MEDIA_EXTENSIONS = {'image/png': '.png', 'image/jpeg': '.jpg'}


def digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


@dataclass(frozen=True)
class FontResource:
    path: Path
    sha256: str
    family: str


@dataclass(frozen=True)
class FontBinding:
    manifest_revision: str
    path: Path
    sha256: str
    family: str
    additional_fonts: tuple[FontResource, ...] = ()


@dataclass(frozen=True)
class ConverterBinding:
    revision: str
    executable: Path
    version: str
    timeout_seconds: int
    sha256: str


@dataclass(frozen=True, repr=False)
class FixedAsset:
    artifact_id: str
    sha256: str
    media_type: str
    content: bytes


@dataclass(frozen=True)
class ExportInput:
    tenant_id: str
    project_id: str
    document_revision: str
    attempt_id: str
    template_version: str
    body_sha256: str
    template_sha256: str
    purpose: Literal['FORMAL_BODY', 'REVIEW_PACKAGE', 'INTERNAL_EXCHANGE']
    input_data_kind: Literal['SIMULATION', 'REGISTERED']
    font: FontBinding
    converter: ConverterBinding


@dataclass(frozen=True)
class ExportFile:
    format: str
    path: Path
    sha256: str
    size_bytes: int
    input_sha256: str
    verified_node_ids: tuple[str, ...]
    required_text: tuple[str, ...]
    plan_sha256: str
    channel_check: str = 'PENDING_REVIEW'
    visual_check: str = 'NOT_RUN'


class ExportRenderer:
    def __init__(self, catalogue: SchemaCatalogue):
        self.bindings = InterfaceBindings(catalogue)

    def _prepare(self, plan: ExportInput, body_bytes: bytes, assets: tuple[FixedAsset, ...]):
        for value in (plan.tenant_id, plan.project_id, plan.document_revision, plan.attempt_id,
                plan.template_version, plan.font.manifest_revision, plan.converter.revision):
            uuid_value(value)
        if plan.purpose not in {'FORMAL_BODY', 'REVIEW_PACKAGE', 'INTERNAL_EXCHANGE'}:
            raise ValueError('导出用途无效')
        if plan.input_data_kind not in {'SIMULATION', 'REGISTERED'}:
            raise ValueError('导出资料身份无效')
        if digest(body_bytes) != plan.body_sha256:
            raise ValueError('完整正文摘要不一致')
        bound = self.bindings.bind(json.loads(body_bytes), 'DocumentBody')
        if bound.canonical_bytes != body_bytes:
            raise ValueError('正文必须使用固定规范字节')
        body = bound.read()
        if body.template_version != plan.template_version:
            raise ValueError('正文模板版本不一致')
        nodes = {node.node_id: node for node in body.nodes}
        if len(nodes) != len(body.nodes) or len(nodes) > 100000:
            raise ValueError('正文节点身份或容量无效')
        roots = [node for node in body.nodes if node.parent_id is None]
        if len(roots) != 1 or roots[0].content.kind != 'DOCUMENT':
            raise ValueError('正文目录身份无效')
        children = defaultdict(list)
        positions = set()
        for node in body.nodes:
            position = (node.parent_id, node.ordinal)
            if position in positions:
                raise ValueError('正文节点顺序重复')
            positions.add(position)
            if node.parent_id is not None:
                if node.parent_id not in nodes or nodes[node.parent_id].content.kind not in {'DOCUMENT', 'SECTION', 'PARAGRAPH', 'TABLE'}:
                    raise ValueError('正文父节点无效')
                children[node.parent_id].append(node)
        for values in children.values():
            values.sort(key=lambda node: node.ordinal)
        reached = set()
        pending = [roots[0]]
        while pending:
            current = pending.pop()
            if current.node_id in reached:
                raise ValueError('正文节点重复或存在循环')
            reached.add(current.node_id)
            pending.extend(children[current.node_id])
        if reached != nodes.keys():
            raise ValueError('正文包含无法读取的节点')
        fields = {field.field_id: field for field in body.fields}
        links = {link.evidence_link_id: link for link in body.evidence_links}
        if len(fields) != len(body.fields) or len(links) != len(body.evidence_links):
            raise ValueError('正文字段或证据身份重复')
        for field in body.fields:
            if field.node_id not in nodes or field.field_id not in nodes[field.node_id].field_ids:
                raise ValueError('正文字段归属无效')
            for link_id in field.evidence_link_ids:
                if link_id not in links or field.node_id not in links[link_id].node_ids:
                    raise ValueError('字段证据范围不一致')
        for link in body.evidence_links:
            evidence = link.evidence
            if evidence.tenant_id != plan.tenant_id or evidence.project_id != plan.project_id:
                raise ValueError('正文证据企业或项目不一致')
            for node_id in link.node_ids:
                if node_id not in nodes or link.evidence_link_id not in nodes[node_id].evidence_link_ids:
                    raise ValueError('正文证据节点范围不一致')
        for node in body.nodes:
            if any(field_id not in fields or fields[field_id].node_id != node.node_id for field_id in node.field_ids) \
                    or any(link_id not in links or node.node_id not in links[link_id].node_ids for link_id in node.evidence_link_ids):
                raise ValueError('正文引用关系不完整')
        targets = set()
        for usage in body.fact_usages:
            if not usage.node_ids and not usage.field_ids:
                raise ValueError('事实采用位置不能为空')
            for target_id, values in ((node_id, nodes) for node_id in usage.node_ids):
                key = (usage.fact_version_id, 'NODE', target_id)
                if target_id not in values or key in targets:
                    raise ValueError('事实节点位置无效')
                targets.add(key)
            for field_id in usage.field_ids:
                key = (usage.fact_version_id, 'FIELD', field_id)
                if field_id not in fields or key in targets:
                    raise ValueError('事实字段位置无效')
                targets.add(key)
        for node in body.nodes:
            if node.content.kind != 'TABLE':
                continue
            table = node.content
            occupied, selected = set(), set()
            for cell in table.cells:
                row = table.row_ids.index(cell.row_id)
                column = table.column_ids.index(cell.column_id)
                if row + cell.row_span > len(table.row_ids) or column + cell.column_span > len(table.column_ids):
                    raise ValueError('表格合并区域无效')
                for row_index in range(row, row + cell.row_span):
                    for column_index in range(column, column + cell.column_span):
                        if (row_index, column_index) in occupied:
                            raise ValueError('表格单元格区域重复')
                        occupied.add((row_index, column_index))
                for node_id in cell.node_ids:
                    if node_id in selected or node_id not in nodes or nodes[node_id].parent_id != node.node_id:
                        raise ValueError('表格正文归属无效')
                    selected.add(node_id)
            if len(occupied) != len(table.row_ids) * len(table.column_ids) or selected != {item.node_id for item in children[node.node_id]}:
                raise ValueError('表格范围不完整')
        images = {}
        for asset in assets:
            uuid_value(asset.artifact_id)
            if asset.artifact_id in images or asset.media_type not in MEDIA_EXTENSIONS or digest(asset.content) != asset.sha256:
                raise ValueError('固定图片身份或摘要无效')
            with Image.open(BytesIO(asset.content)) as image:
                if image.format != {'image/png': 'PNG', 'image/jpeg': 'JPEG'}[asset.media_type] or getattr(image, 'n_frames', 1) != 1:
                    raise ValueError('图片实际格式无效')
                image.verify()
            images[asset.artifact_id] = asset
        requested_assets = set()
        for node in body.nodes:
            if node.content.kind == 'IMAGE':
                asset = images.get(node.content.artifact_id)
                if asset is None or asset.sha256 != node.content.sha256:
                    raise ValueError('正文图片未完整提供')
                requested_assets.add(asset.artifact_id)
        if requested_assets != images.keys():
            raise ValueError('导出图片清单与正文不一致')
        strings = self._texts(body)
        if plan.input_data_kind == 'SIMULATION':
            strings.insert(0, SIMULATION_LABEL)
        self._font(plan.font, strings)
        identity = digest(rfc8785.dumps({
            'tenant_id': plan.tenant_id, 'project_id': plan.project_id, 'document_revision': plan.document_revision,
            'attempt_id': plan.attempt_id, 'body_sha256': plan.body_sha256, 'template_version': plan.template_version,
            'template_sha256': plan.template_sha256, 'font_manifest_revision': plan.font.manifest_revision,
            'font_sha256': plan.font.sha256, 'font_family': plan.font.family, 'converter_revision': plan.converter.revision,
            'additional_fonts': [{'sha256': font.sha256, 'family': font.family} for font in plan.font.additional_fonts],
            'converter_sha256': plan.converter.sha256,
            'converter_version': plan.converter.version, 'purpose': plan.purpose, 'input_data_kind': plan.input_data_kind,
            'assets': [{'artifact_id': asset.artifact_id, 'sha256': asset.sha256, 'media_type': asset.media_type} for asset in assets],
        }))
        return body, nodes, children, roots[0], images, tuple(strings), identity

    @staticmethod
    def _font(font: FontBinding, texts):
        if not font.path.is_file() or digest(font.path.read_bytes()) != font.sha256 or not font.family:
            raise ValueError('字体完整文件摘要不一致')
        with TTFont(font.path, lazy=False) as actual:
            families = {value.toUnicode() for value in actual['name'].names if value.nameID in {1, 16}}
            if font.family not in families or actual['OS/2'].fsType & 2:
                raise ValueError('字体名称或嵌入许可无效')
            available = actual.getBestCmap()
            for text in texts:
                if any(not character.isspace() and ord(character) not in available for character in text):
                    raise ValueError('字体未覆盖完整正文字符')
        for resource in font.additional_fonts:
            if not resource.path.is_file() or digest(resource.path.read_bytes()) != resource.sha256:
                raise ValueError('字体清单制品摘要不一致')
            with TTFont(resource.path, lazy=False) as actual:
                families = {value.toUnicode() for value in actual['name'].names if value.nameID in {1, 16}}
                if resource.family not in families or actual['OS/2'].fsType & 2:
                    raise ValueError('字体清单名称或嵌入许可无效')

    @staticmethod
    def _texts(body: DocumentBody):
        values = []
        for node in body.nodes:
            content = node.content
            if content.kind in {'DOCUMENT', 'SECTION', 'PARAGRAPH'} and content.title:
                values.append(content.title)
            elif content.kind == 'TEXT' and content.text:
                values.append(content.text)
            elif content.kind in {'IMAGE', 'TABLE'} and content.caption:
                values.append(content.caption)
            elif content.kind == 'FORMULA' and content.description:
                values.append(content.description)
        return values

    @staticmethod
    def _output(directory: Path):
        absolute = directory.parent.resolve(strict=True) / directory.name
        if absolute.is_relative_to(Path('/tmp')) or absolute.is_relative_to(Path('/private/tmp')):
            raise ValueError('转换目录必须位于项目工作目录')
        absolute.mkdir(mode=0o700, parents=False, exist_ok=False)
        return absolute

    @staticmethod
    def _anchor(paragraph, node_id, number):
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(number))
        start.set(qn('w:name'), 'medical_' + node_id.replace('-', '_'))
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(number))
        paragraph._p.append(start)
        paragraph._p.append(end)

    def docx(self, plan: ExportInput, body_bytes: bytes, template_bytes: bytes, assets: tuple[FixedAsset, ...], directory: Path) -> ExportFile:
        body, nodes, children, root, images, texts, identity = self._prepare(plan, body_bytes, assets)
        if digest(template_bytes) != plan.template_sha256:
            raise ValueError('格式模板完整摘要不一致')
        document = Document(BytesIO(template_bytes))
        if any(paragraph.text for paragraph in document.paragraphs) or document.tables or document.inline_shapes:
            raise ValueError('格式模板必须提供空白正文区域')
        for part in document.part.package.parts:
            if any(relation.is_external for relation in part.rels.values()):
                raise ValueError('格式模板包含外部资源')
        for style_name in ('Normal', 'Title', *('Heading ' + str(level) for level in range(1, 10))):
            style = document.styles[style_name]
            style.font.name = plan.font.family
            properties = style.element.get_or_add_rPr()
            fonts = properties.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                properties.append(fonts)
            fonts.set(qn('w:eastAsia'), plan.font.family)
            fonts.set(qn('w:cs'), plan.font.family)
            for attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                fonts.attrib.pop(qn('w:' + attribute), None)
        math_properties = document.settings.element.find(qn('m:mathPr'))
        if math_properties is None:
            math_properties = OxmlElement('m:mathPr')
            document.settings.element.append(math_properties)
        math_font = math_properties.find(qn('m:mathFont'))
        if math_font is None:
            math_font = OxmlElement('m:mathFont')
            math_properties.append(math_font)
        math_font.set(qn('m:val'), plan.font.family)
        if plan.input_data_kind == 'SIMULATION':
            document.add_paragraph(SIMULATION_LABEL)
        processed = []
        def render(node, target, level, paragraph=None):
            processed.append(node.node_id)
            content = node.content
            if content.kind in {'DOCUMENT', 'SECTION', 'PARAGRAPH'}:
                if content.kind == 'DOCUMENT':
                    current = target.add_paragraph(content.title, style='Title')
                elif content.kind == 'SECTION':
                    if level > 9:
                        raise ValueError('章节层级超过固定格式支持范围')
                    current = target.add_paragraph(content.title, style='Heading ' + str(level))
                else:
                    current = target.add_paragraph(content.title)
                self._anchor(current, node.node_id, len(processed))
                for child in children[node.node_id]:
                    render(child, target, level + (content.kind == 'SECTION'), current if content.kind == 'PARAGRAPH' else None)
            elif content.kind == 'TEXT':
                current = paragraph if paragraph is not None else target.add_paragraph()
                self._anchor(current, node.node_id, len(processed))
                run = current.add_run(content.text)
                run.bold = 'BOLD' in content.marks
                run.italic = 'ITALIC' in content.marks
                run.underline = 'UNDERLINE' in content.marks
                run.font.superscript = 'SUPERSCRIPT' in content.marks
                run.font.subscript = 'SUBSCRIPT' in content.marks
            elif content.kind == 'TABLE':
                caption = target.add_paragraph(content.caption)
                self._anchor(caption, node.node_id, len(processed))
                table = target.add_table(rows=len(content.row_ids), cols=len(content.column_ids))
                for cell in content.cells:
                    row = content.row_ids.index(cell.row_id)
                    column = content.column_ids.index(cell.column_id)
                    selected = table.cell(row, column)
                    if cell.row_span > 1 or cell.column_span > 1:
                        selected = selected.merge(table.cell(row + cell.row_span - 1, column + cell.column_span - 1))
                    for node_id in cell.node_ids:
                        render(nodes[node_id], selected, level)
            elif content.kind == 'FORMULA':
                current = paragraph if paragraph is not None else target.add_paragraph()
                self._anchor(current, node.node_id, len(processed))
                omml = mathml_to_omml(latex_to_mathml(content.latex))
                math_element = parse_xml('<medical ' + nsdecls('m', 'w') + '>' + omml + '</medical>')[0]
                for math_run in math_element.findall('.//' + qn('m:r')):
                    run_properties = OxmlElement('w:rPr')
                    run_fonts = OxmlElement('w:rFonts')
                    for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                        run_fonts.set(qn('w:' + attribute), plan.font.family)
                    run_properties.append(run_fonts)
                    math_run.insert(1, run_properties)
                current._p.append(math_element)
                if content.description:
                    target.add_paragraph(content.description)
            elif content.kind == 'IMAGE':
                current = paragraph if paragraph is not None else target.add_paragraph()
                self._anchor(current, node.node_id, len(processed))
                section = document.sections[0]
                width = section.page_width - section.left_margin - section.right_margin
                shape = current.add_run().add_picture(BytesIO(images[content.artifact_id].content), width=width)
                shape._inline.docPr.set('descr', content.alt_text)
                if content.caption:
                    target.add_paragraph(content.caption)
        render(root, document, 1)
        if set(processed) != nodes.keys() or len(processed) != len(nodes):
            raise ValueError('正文生成范围不完整')
        output = self._output(directory)
        path = output / 'body.docx'
        document.save(path)
        actual = Document(path)
        actual_ids = tuple(value.get(qn('w:name')) for value in actual.element.xpath('.//w:bookmarkStart'))
        if actual_ids != tuple('medical_' + node_id.replace('-', '_') for node_id in processed):
            raise ValueError('DOCX节点重读与固定正文不一致')
        actual_text = ''.join(value.text or '' for value in actual.element.xpath('.//w:t'))
        self._require_text(texts, actual_text)
        for kind, query in (('TABLE', './/w:tbl'), ('FORMULA', './/m:oMath'), ('IMAGE', './/wp:inline')):
            if sum(node.content.kind == kind for node in body.nodes) != len(actual.element.xpath(query)):
                raise ValueError('DOCX结构重读与正文不一致')
        self._projections(plan, body, output)
        return self._file('DOCX', path, identity, tuple(processed), texts, self._plan_digest(plan))

    @staticmethod
    def _require_text(required, actual):
        compact = ''.join(actual.split())
        for value in required:
            if ''.join(value.split()) not in compact:
                raise ValueError('成品未完整保留固定正文')

    @staticmethod
    def _file(format_name, path, identity, nodes, texts, plan_sha256):
        content = path.read_bytes()
        if not content:
            raise ValueError('成品文件不能为空')
        return ExportFile(format_name, path, digest(content), len(content), identity, nodes, texts, plan_sha256)

    @staticmethod
    def _plan_digest(plan):
        return digest(rfc8785.dumps({'tenant_id': plan.tenant_id, 'project_id': plan.project_id,
            'document_revision': plan.document_revision, 'attempt_id': plan.attempt_id, 'template_version': plan.template_version,
            'body_sha256': plan.body_sha256, 'template_sha256': plan.template_sha256, 'purpose': plan.purpose,
            'input_data_kind': plan.input_data_kind, 'font_manifest_revision': plan.font.manifest_revision,
            'font_sha256': plan.font.sha256, 'font_family': plan.font.family,
            'additional_fonts': [{'sha256': font.sha256, 'family': font.family} for font in plan.font.additional_fonts],
            'converter_revision': plan.converter.revision, 'converter_version': plan.converter.version,
            'converter_sha256': plan.converter.sha256, 'converter_timeout_seconds': plan.converter.timeout_seconds}))

    @staticmethod
    def _projections(plan, body, directory):
        if plan.purpose == 'FORMAL_BODY':
            return
        with (directory / 'internal-projections.json').open('xb') as stream:
            stream.write(rfc8785.dumps({'schema_version': '1.0.0', 'document_revision': plan.document_revision,
                'body_sha256': plan.body_sha256, 'input_data_kind': plan.input_data_kind,
                'fields': [value.model_dump(mode='json') for value in body.fields],
                'evidence_links': [value.model_dump(mode='json') for value in body.evidence_links],
                'fact_usages': [value.model_dump(mode='json') for value in body.fact_usages]}))

    def pdf(self, plan: ExportInput, document: ExportFile, directory: Path) -> ExportFile:
        converter = plan.converter
        if document.format != 'DOCX' or digest(document.path.read_bytes()) != document.sha256 or not document.verified_node_ids \
                or document.plan_sha256 != self._plan_digest(plan):
            raise ValueError('PDF转换必须读取原已核验DOCX')
        if converter.timeout_seconds < 1 or converter.timeout_seconds > 300 or not converter.executable.is_file() \
                or digest(converter.executable.read_bytes()) != converter.sha256:
            raise ValueError('转换程序与执行期限无效')
        version = subprocess.run([str(converter.executable), '--version'], check=True, stdout=subprocess.PIPE,
            stderr=subprocess.PIPE, timeout=converter.timeout_seconds).stdout.decode('utf-8').strip()
        if version != converter.version:
            raise ValueError('转换程序版本与固定清单不一致')
        self._font(plan.font, document.required_text)
        output = self._output(directory)
        profile = output / 'libreoffice-profile'
        profile.mkdir(mode=0o700)
        environment = os.environ.copy()
        environment['TMPDIR'] = str(output)
        # 转换进程使用固定字体清单和独立缓存；字体准入与实际使用必须一致。
        font_directory = output / 'fonts'
        font_directory.mkdir(mode=0o700)
        postscript_names = set()
        for number, resource in enumerate((plan.font, *plan.font.additional_fonts)):
            fixed_font = font_directory / (str(number) + resource.path.suffix)
            font_bytes = resource.path.read_bytes()
            if digest(font_bytes) != resource.sha256:
                raise ValueError('转换字体文件发生变化')
            with TTFont(BytesIO(font_bytes)) as actual_font:
                postscript_names.update(value.toUnicode() for value in actual_font['name'].names if value.nameID == 6)
            with fixed_font.open('xb') as stream:
                stream.write(font_bytes)
        font_configuration = etree.Element('fontconfig')
        etree.SubElement(font_configuration, 'dir').text = str(font_directory)
        etree.SubElement(font_configuration, 'cachedir').text = str(output / 'font-cache')
        config_path = output / 'fonts.conf'
        with config_path.open('xb') as stream:
            stream.write(etree.tostring(font_configuration, encoding='utf-8', xml_declaration=True,
                doctype='<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">'))
        environment['FONTCONFIG_FILE'] = str(config_path)
        environment['FONTCONFIG_PATH'] = str(output)
        command = [str(converter.executable), '-env:UserInstallation=' + profile.as_uri(), '--headless', '--nologo',
            '--nodefault', '--norestore', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(output), str(document.path)]
        process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, cwd=output,
            env=environment, start_new_session=True)
        try:
            process.communicate(timeout=converter.timeout_seconds)
            if process.returncode != 0:
                raise subprocess.CalledProcessError(process.returncode, command)
        finally:
            if process.poll() is None:
                os.killpg(process.pid, signal.SIGKILL)
                process.wait()
        if digest(document.path.read_bytes()) != document.sha256:
            raise ValueError('转换期间原DOCX发生变化')
        path = output / (document.path.stem + '.pdf')
        actual = PdfReader(path, strict=True)
        if actual.is_encrypted or not actual.pages:
            raise ValueError('PDF成品结构无效')
        for page in actual.pages:
            for font in page['/Resources']['/Font'].values():
                name = str(font.get_object()['/BaseFont'])
                if not any(name == '/' + expected or name.endswith('+' + expected) for expected in postscript_names):
                    raise ValueError('PDF实际字体不在固定清单中')
        self._require_text(document.required_text, ''.join(page.extract_text() for page in actual.pages))
        return self._file('PDF', path, document.input_sha256, document.verified_node_ids, document.required_text, document.plan_sha256)

    def markdown(self, plan: ExportInput, body_bytes: bytes, assets: tuple[FixedAsset, ...], directory: Path) -> ExportFile:
        if plan.purpose != 'INTERNAL_EXCHANGE':
            raise ValueError('Markdown只用于内部交换')
        body, nodes, children, root, images, texts, identity = self._prepare(plan, body_bytes, assets)
        article = etree.Element('article')
        if plan.input_data_kind == 'SIMULATION':
            etree.SubElement(article, 'p').text = SIMULATION_LABEL
        processed = []
        def render(node, target, level):
            processed.append(node.node_id)
            content = node.content
            if content.kind in {'DOCUMENT', 'SECTION', 'PARAGRAPH'}:
                current = etree.SubElement(target, 'section' if content.kind in {'DOCUMENT', 'SECTION'} else 'p', id=node.node_id)
                if content.title:
                    if level > 6 and content.kind != 'PARAGRAPH':
                        raise ValueError('Markdown章节层级超出支持范围')
                    etree.SubElement(current, 'h' + str(level) if content.kind != 'PARAGRAPH' else 'span').text = content.title
                for child in children[node.node_id]:
                    render(child, current, level + (content.kind == 'SECTION'))
            elif content.kind == 'TEXT':
                current = etree.SubElement(target, 'span', id=node.node_id)
                for mark in content.marks:
                    current = etree.SubElement(current, {'BOLD': 'strong', 'ITALIC': 'em', 'UNDERLINE': 'u', 'SUPERSCRIPT': 'sup', 'SUBSCRIPT': 'sub'}[mark])
                current.text = content.text
            elif content.kind == 'TABLE':
                table = etree.SubElement(target, 'table', id=node.node_id)
                etree.SubElement(table, 'caption').text = content.caption
                rows = [etree.SubElement(table, 'tr') for _ in content.row_ids]
                for cell in sorted(content.cells, key=lambda value: (content.row_ids.index(value.row_id), content.column_ids.index(value.column_id))):
                    current = etree.SubElement(rows[content.row_ids.index(cell.row_id)], 'td', rowspan=str(cell.row_span), colspan=str(cell.column_span))
                    for node_id in cell.node_ids:
                        render(nodes[node_id], current, level)
            elif content.kind == 'FORMULA':
                current = etree.SubElement(target, 'figure', id=node.node_id)
                current.append(etree.fromstring(latex_to_mathml(content.latex).encode('utf-8'), etree.XMLParser(resolve_entities=False, no_network=True)))
                etree.SubElement(current, 'pre', attrib={'data-formula-latex': 'true'}).text = content.latex
                etree.SubElement(current, 'figcaption').text = content.description
            elif content.kind == 'IMAGE':
                current = etree.SubElement(target, 'figure', id=node.node_id)
                asset = images[content.artifact_id]
                etree.SubElement(current, 'img', src=asset.artifact_id + MEDIA_EXTENSIONS[asset.media_type], alt=content.alt_text)
                etree.SubElement(current, 'figcaption').text = content.caption
        render(root, article, 1)
        if set(processed) != nodes.keys() or len(processed) != len(nodes):
            raise ValueError('Markdown生成范围不完整')
        value = etree.tostring(article, encoding='utf-8', method='html')
        reread = etree.fromstring(value, etree.HTMLParser(encoding='utf-8', no_network=True))
        self._require_text(texts, ''.join(reread.itertext()))
        output = self._output(directory)
        path = output / 'body.md'
        with path.open('xb') as stream:
            stream.write(value + b'\n')
        for asset in assets:
            with (output / (asset.artifact_id + MEDIA_EXTENSIONS[asset.media_type])).open('xb') as stream:
                stream.write(asset.content)
        self._projections(plan, body, output)
        return self._file('MARKDOWN', path, identity, tuple(processed), texts, self._plan_digest(plan))
